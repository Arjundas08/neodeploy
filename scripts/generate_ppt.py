"""
NeoDeploy - PowerPoint Presentation Generator
Creates a professional PPT for project presentation
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 aspect ratio
prs.slide_height = Inches(7.5)

# Color scheme
DARK_BG = RGBColor(15, 23, 42)       # Dark blue
ACCENT = RGBColor(59, 130, 246)       # Blue
SUCCESS = RGBColor(34, 197, 94)       # Green
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(148, 163, 184)
DARK_BAR = RGBColor(30, 41, 59)

def add_title_slide(title, subtitle):
    """Add a title slide"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(1))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(24)
    p.font.color.rgb = GRAY
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(title, bullets, emoji=""):
    """Add a content slide with bullets"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Title bar
    title_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
    title_bar.fill.solid()
    title_bar.fill.fore_color.rgb = DARK_BAR
    title_bar.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"{emoji} {title}" if emoji else title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Content
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.5), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {bullet}"
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE
        p.space_after = Pt(18)
    
    return slide

def add_architecture_slide():
    """Add architecture diagram slide"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "🏗️ System Architecture"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Architecture boxes
    boxes = [
        ("GitHub", 0.5, 3),
        ("Jenkins", 3, 3),
        ("Docker", 5.5, 3),
        ("Deploy", 8, 3),
        ("Monitor", 10.5, 3),
    ]
    
    for text, x, y in boxes:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2), Inches(1))
        box.fill.solid()
        box.fill.fore_color.rgb = ACCENT
        box.line.fill.background()
        
        tf = box.text_frame
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.anchor = MSO_ANCHOR.MIDDLE
    
    # Arrows (as text)
    arrows_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.3), Inches(12), Inches(0.5))
    tf = arrows_box.text_frame
    p = tf.paragraphs[0]
    p.text = "     Push  →      Build/Test  →    Containerize  →    Deploy  →    Alert"
    p.font.size = Pt(16)
    p.font.color.rgb = GRAY
    
    # Telegram box at bottom
    tg_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.5), Inches(5.5), Inches(2.5), Inches(0.8))
    tg_box.fill.solid()
    tg_box.fill.fore_color.rgb = SUCCESS
    tg_box.line.fill.background()
    tf = tg_box.text_frame
    tf.paragraphs[0].text = "📱 Telegram"
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.anchor = MSO_ANCHOR.MIDDLE
    
    return slide

def add_pipeline_slide():
    """Add Jenkins pipeline stages slide"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "⚡ Jenkins Pipeline Stages"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Pipeline stages
    stages = [
        ("1. BUILD", "Maven compiles Java code", ACCENT),
        ("2. TEST", "11 JUnit tests run", ACCENT),
        ("3. DOCKER", "Container image created", ACCENT),
        ("4. SECURITY", "Trivy scans for CVEs", ACCENT),
        ("5. DEPLOY", "Container deployed", ACCENT),
        ("6. HEALTH", "Endpoint verified", SUCCESS),
    ]
    
    y_pos = 1.5
    for stage, desc, color in stages:
        # Stage box
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(y_pos), Inches(2.5), Inches(0.7))
        box.fill.solid()
        box.fill.fore_color.rgb = color
        box.line.fill.background()
        tf = box.text_frame
        tf.paragraphs[0].text = stage
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.anchor = MSO_ANCHOR.MIDDLE
        
        # Description
        desc_box = slide.shapes.add_textbox(Inches(4), Inches(y_pos + 0.15), Inches(8), Inches(0.5))
        tf = desc_box.text_frame
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(20)
        tf.paragraphs[0].font.color.rgb = WHITE
        
        y_pos += 0.9
    
    return slide

def add_tech_stack_slide():
    """Add technology stack slide"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "🛠️ Technology Stack"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Tech grid - Left column
    left_tech = [
        "☕ Java 17",
        "🍃 Spring Boot 3.2.5",
        "🔧 Maven",
        "🧪 JUnit 5",
        "📊 JaCoCo",
    ]
    
    # Right column
    right_tech = [
        "🐳 Docker",
        "🔄 Jenkins",
        "📦 Ansible",
        "🔒 Trivy Security",
        "📱 Telegram Bot",
    ]
    
    # Left column
    left_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(5))
    tf = left_box.text_frame
    for i, tech in enumerate(left_tech):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = tech
        p.font.size = Pt(28)
        p.font.color.rgb = WHITE
        p.space_after = Pt(20)
    
    # Right column
    right_box = slide.shapes.add_textbox(Inches(7), Inches(1.5), Inches(5), Inches(5))
    tf = right_box.text_frame
    for i, tech in enumerate(right_tech):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = tech
        p.font.size = Pt(28)
        p.font.color.rgb = WHITE
        p.space_after = Pt(20)
    
    return slide

def add_demo_slide():
    """Add live demo slide"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "🎬 LIVE DEMO"
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER
    
    # URLs
    urls_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(12.333), Inches(2))
    tf = urls_box.text_frame
    
    p = tf.paragraphs[0]
    p.text = "Dashboard: http://localhost:9090"
    p.font.size = Pt(24)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    p = tf.add_paragraph()
    p.text = "Jenkins: http://localhost:8081"
    p.font.size = Pt(24)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_conclusion_slide():
    """Add conclusion slide"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "✅ Key Takeaways"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Takeaways
    takeaways = [
        "✓ Fully automated CI/CD pipeline",
        "✓ Real Jenkins integration (not simulation)",
        "✓ Docker containerization",
        "✓ Security scanning with Trivy",
        "✓ Infrastructure as Code (Ansible)",
        "✓ Self-healing capabilities",
        "✓ Real-time monitoring dashboard",
        "✓ Instant Telegram notifications",
    ]
    
    content_box = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(11), Inches(5))
    tf = content_box.text_frame
    
    for i, point in enumerate(takeaways):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = point
        p.font.size = Pt(26)
        p.font.color.rgb = SUCCESS
        p.space_after = Pt(12)
    
    return slide

def add_thank_you_slide():
    """Add thank you slide"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    background.fill.solid()
    background.fill.fore_color.rgb = DARK_BG
    background.line.fill.background()
    
    # Thank you text
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Thank You! 🙏"
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Questions
    q_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(12.333), Inches(1))
    tf = q_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Questions?"
    p.font.size = Pt(32)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER
    
    return slide

# ==========================================
# Generate Presentation
# ==========================================

print("🎨 Generating NeoDeploy Presentation...")

# Slide 1: Title
add_title_slide("NeoDeploy", "CI/CD DevOps Pipeline with Self-Healing & Real-Time Monitoring")

# Slide 2: Problem Statement
add_content_slide("Problem Statement", [
    "Manual deployments are slow and error-prone",
    "No visibility into deployment status",
    "Developers waste time on repetitive tasks",
    "No automated testing or security checks",
    "No instant alerts when things break",
], "❓")

# Slide 3: Solution
add_content_slide("Our Solution: NeoDeploy", [
    "Automated CI/CD pipeline with Jenkins",
    "One-click deployments from dashboard",
    "Real-time monitoring and health checks",
    "Automatic security vulnerability scanning",
    "Instant Telegram notifications",
    "Self-healing infrastructure with Ansible",
], "💡")

# Slide 4: Architecture
add_architecture_slide()

# Slide 5: Pipeline Stages
add_pipeline_slide()

# Slide 6: Technology Stack
add_tech_stack_slide()

# Slide 7: Features
add_content_slide("Key Features", [
    "🖥️ Real-time Dashboard - Live deployment monitoring",
    "🔄 Jenkins Integration - Automated build & test",
    "🐳 Docker - Containerized deployments",
    "🔒 Trivy - Security vulnerability scanning",
    "📦 Ansible - Infrastructure automation",
    "🔧 Self-Healing - Auto-restart failed containers",
    "📱 Telegram - Instant build notifications",
], "⭐")

# Slide 8: Live Demo
add_demo_slide()

# Slide 9: Conclusion
add_conclusion_slide()

# Slide 10: Thank You
add_thank_you_slide()

# Save presentation
output_path = r"C:\Users\hp\OneDrive\Desktop\Neo-Deploy\NeoDeploy-Presentation.pptx"
prs.save(output_path)

print(f"✅ Presentation saved to: {output_path}")
print(f"📊 Total slides: {len(prs.slides)}")
